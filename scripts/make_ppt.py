#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成"三比值法 × 贝叶斯网络混合归因"两页答辩 PPT。"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "混合故障归因_答辩PPT.pptx"

# ── 配色（Midnight Executive / 学术深蓝）──
NAVY   = RGBColor(0x1E, 0x2A, 0x4A)   # 主深蓝
BLUE   = RGBColor(0x2E, 0x5A, 0x9E)   # 中蓝
LBLUE  = RGBColor(0xDC, 0xE6, 0xF5)   # 浅蓝底
ICE    = RGBColor(0xCA, 0xDC, 0xFC)   # 冰蓝
ACCENT = RGBColor(0xF2, 0xA6, 0x3B)   # 橙色强调
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
DARK   = RGBColor(0x22, 0x2A, 0x35)   # 正文深灰
GRAY   = RGBColor(0x5A, 0x63, 0x70)

HEAD_FONT = "微软雅黑"
BODY_FONT = "微软雅黑"
MONO_FONT = "Consolas"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_rect(slide, x, y, w, h, fill, line=None, round_=False):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
        x, y, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp


def add_text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             space_after=6, line_spacing=1.05):
    """runs: list of paragraphs; each paragraph = list of (text, size, color, bold, font, italic)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(2)
    tf.margin_top = tf.margin_bottom = Pt(2)
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        for (txt, size, color, bold, font, italic) in para:
            r = p.add_run(); r.text = txt
            r.font.size = Pt(size); r.font.color.rgb = color
            r.font.bold = bold; r.font.name = font; r.font.italic = italic
    return tb


def R(txt, size=14, color=DARK, bold=False, font=BODY_FONT, italic=False):
    return (txt, size, color, bold, font, italic)


def title_bar(slide, title, page):
    add_rect(slide, 0, 0, SW, Inches(1.15), NAVY)
    add_rect(slide, 0, Inches(1.15), SW, Pt(3), ACCENT)
    add_text(slide, Inches(0.55), Inches(0.16), Inches(11.5), Inches(0.85),
             [[R(title, 27, WHITE, True, HEAD_FONT)]], anchor=MSO_ANCHOR.MIDDLE)
    # 页码圆
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, SW - Inches(1.05), Inches(0.3), Inches(0.55), Inches(0.55))
    c.fill.solid(); c.fill.fore_color.rgb = ACCENT; c.line.fill.background(); c.shadow.inherit = False
    add_text(slide, SW - Inches(1.05), Inches(0.3), Inches(0.55), Inches(0.55),
             [[R(page, 18, NAVY, True, HEAD_FONT)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def section_chip(slide, x, y, w, text):
    """小节标题（深蓝条）。"""
    add_rect(slide, x, y, w, Inches(0.4), BLUE, round_=True)
    add_text(slide, x + Inches(0.12), y, w - Inches(0.2), Inches(0.4),
             [[R(text, 14, WHITE, True, HEAD_FONT)]], anchor=MSO_ANCHOR.MIDDLE)


# ════════════════════════════════════════════════════════════
# 第 1 页：方法原理与四步流程
# ════════════════════════════════════════════════════════════
s1 = prs.slides.add_slide(BLANK)
add_rect(s1, 0, 0, SW, SH, WHITE)
title_bar(s1, "三比值法 × 贝叶斯网络混合故障归因（一）原理", "1")

# 左栏：设计动机 + 网络结构
LX = Inches(0.55); LW = Inches(5.7)
section_chip(s1, LX, Inches(1.45), LW, "为什么要“混合”？——设计动机")
add_text(s1, LX, Inches(1.95), LW, Inches(1.7), [
    [R("• 三比值法（DL/T 722 国标）", 14, NAVY, True, BODY_FONT)],
    [R("   可解释、合规，但", 13, DARK, False, BODY_FONT), R("只判定类型、无法量化概率", 13, ACCENT, True, BODY_FONT)],
    [R("• 贝叶斯网络（概率推理）", 14, NAVY, True, BODY_FONT)],
    [R("   可量化故障概率，但", 13, DARK, False, BODY_FONT), R("可解释性差、专家不信任", 13, ACCENT, True, BODY_FONT)],
    [R("→ 本方法将二者融合：", 14, BLUE, True, BODY_FONT), R("以概率为主、规则修正，兼取所长", 13, DARK, False, BODY_FONT)],
], space_after=5, line_spacing=1.1)

section_chip(s1, LX, Inches(3.85), LW, "贝叶斯网络结构（两层有向图）")
# 结构示意：故障层 → 征兆层
fx, fy = LX, Inches(4.45)
add_rect(s1, fx, fy, Inches(2.5), Inches(0.95), LBLUE, line=BLUE, round_=True)
add_text(s1, fx, fy, Inches(2.5), Inches(0.95), [
    [R("故障层 F", 15, NAVY, True, HEAD_FONT)],
    [R("8 类故障", 12, GRAY, False, BODY_FONT)],
    [R("短路/局放/过热…", 11, GRAY, False, BODY_FONT)],
], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=2)

arr = s1.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, fx + Inches(2.55), fy + Inches(0.3), Inches(0.55), Inches(0.35))
arr.fill.solid(); arr.fill.fore_color.rgb = ACCENT; arr.line.fill.background(); arr.shadow.inherit = False

add_rect(s1, fx + Inches(3.2), fy, Inches(2.5), Inches(0.95), LBLUE, line=BLUE, round_=True)
add_text(s1, fx + Inches(3.2), fy, Inches(2.5), Inches(0.95), [
    [R("征兆层 S", 15, NAVY, True, HEAD_FONT)],
    [R("14 类征兆", 12, GRAY, False, BODY_FONT)],
    [R("气体超标/油温升高…", 11, GRAY, False, BODY_FONT)],
], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=2)
add_text(s1, LX, fy + Inches(1.05), LW, Inches(0.5),
         [[R("“朴素”条件独立假设：给定故障时各征兆独立 → 似然可连乘", 11.5, GRAY, False, BODY_FONT, True)]])

# 右栏：四步流程（纵向流程卡）
RX = Inches(6.65); RW = Inches(6.15)
section_chip(s1, RX, Inches(1.45), RW, "四步推理流程")
steps = [
    ("①", "DGA 三比值法编码", "将 C₂H₂/C₂H₄、CH₄/H₂、C₂H₄/C₂H₆ 离散编码为 0/1/2，匹配表7故障判据，得“规则证据”"),
    ("②", "朴素贝叶斯后验推理", "由先验 P(F) 与条件概率表 CPT，计算各故障后验概率并归一化"),
    ("③", "规则—概率融合", "0.7×贝叶斯 + 0.3×规则置信，对命中故障加权修正"),
    ("④", "排序 · 分级 · 报告", "按融合概率排序，映射高/中/低风险，输出可解释诊断报告"),
]
sy = Inches(2.0)
for num, head, desc in steps:
    add_rect(s1, RX, sy, RW, Inches(1.12), LBLUE, round_=True)
    # 序号圆
    oc = s1.shapes.add_shape(MSO_SHAPE.OVAL, RX + Inches(0.18), sy + Inches(0.28), Inches(0.55), Inches(0.55))
    oc.fill.solid(); oc.fill.fore_color.rgb = NAVY; oc.line.fill.background(); oc.shadow.inherit = False
    add_text(s1, RX + Inches(0.18), sy + Inches(0.28), Inches(0.55), Inches(0.55),
             [[R(num, 20, WHITE, True, HEAD_FONT)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s1, RX + Inches(0.95), sy + Inches(0.1), RW - Inches(1.1), Inches(0.95), [
        [R(head, 15, NAVY, True, HEAD_FONT)],
        [R(desc, 12, DARK, False, BODY_FONT)],
    ], space_after=3, line_spacing=1.05)
    sy += Inches(1.22)

# ════════════════════════════════════════════════════════════
# 第 2 页：核心公式与算例
# ════════════════════════════════════════════════════════════
s2 = prs.slides.add_slide(BLANK)
add_rect(s2, 0, 0, SW, SH, WHITE)
title_bar(s2, "三比值法 × 贝叶斯网络混合故障归因（二）公式", "2")

# 左栏：三比值编码表 + 公式
LX = Inches(0.55); LW = Inches(6.0)
section_chip(s2, LX, Inches(1.45), LW, "第①步：三比值法编码规则（DL/T 722 表6）")
# 编码表
tbl = s2.shapes.add_table(4, 2, LX, Inches(1.95), LW, Inches(1.5)).table
tbl.columns[0].width = Inches(3.1); tbl.columns[1].width = Inches(2.9)
hdr = [("特征比值", "编码规则"),
       ("C₂H₂ / C₂H₄", "<0.1→0 ; [0.1,3)→1 ; ≥3→2"),
       ("CH₄ / H₂", "<0.1→1 ; [0.1,1)→0 ; ≥1→2"),
       ("C₂H₄ / C₂H₆", "<1→0 ; [1,3)→1 ; ≥3→2")]
for i, (a, b) in enumerate(hdr):
    for j, txt in enumerate((a, b)):
        cell = tbl.cell(i, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY if i == 0 else (LBLUE if i % 2 else WHITE)
        tf = cell.text_frame; tf.word_wrap = True
        tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = txt
        r.font.size = Pt(12.5 if i == 0 else 12)
        r.font.bold = (i == 0)
        r.font.color.rgb = WHITE if i == 0 else DARK
        r.font.name = BODY_FONT

section_chip(s2, LX, Inches(3.7), LW, "第②步：朴素贝叶斯后验推理")
add_rect(s2, LX, Inches(4.18), LW, Inches(0.95), LBLUE, round_=True)
add_text(s2, LX, Inches(4.18), LW, Inches(0.95),
         [[R("P(Fᵢ | E) ∝ P(Fᵢ) · ∏  P(eⱼ | Fᵢ)", 19, NAVY, True, MONO_FONT)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_text(s2, LX, Inches(5.2), LW, Inches(1.5), [
    [R("• P(Fᵢ)", 13, BLUE, True, BODY_FONT), R("：故障先验（常见程度，如绝缘老化0.18）", 12.5, DARK, False, BODY_FONT)],
    [R("• P(eⱼ|Fᵢ)", 13, BLUE, True, BODY_FONT), R("：条件概率表 CPT（故障发生时征兆出现概率）", 12.5, DARK, False, BODY_FONT)],
    [R("• 仅对观测到的征兆 eⱼ 连乘，最后归一化为概率分布", 12.5, DARK, False, BODY_FONT)],
], space_after=5, line_spacing=1.1)

# 右栏：融合公式 + 算例
RX = Inches(6.85); RW = Inches(5.95)
section_chip(s2, RX, Inches(1.45), RW, "第③步：规则—概率融合")
add_rect(s2, RX, Inches(1.95), RW, Inches(1.3), LBLUE, round_=True)
add_text(s2, RX, Inches(2.05), RW, Inches(1.2), [
    [R("Pfused(Fᵢ) =", 16, NAVY, True, MONO_FONT)],
    [R("0.7·Pbayes + 0.3·Prule", 17, NAVY, True, MONO_FONT), R("（命中规则）", 12, GRAY, False, BODY_FONT)],
    [R("Pbayes", 16, NAVY, True, MONO_FONT), R("（未命中规则）", 12, GRAY, False, BODY_FONT)],
], align=PP_ALIGN.CENTER, space_after=3, line_spacing=1.05)
add_text(s2, RX, Inches(3.3), RW, Inches(0.4),
         [[R("以数据驱动概率为主(70%)，国标规则经验修正(30%)", 11.5, GRAY, False, BODY_FONT, True)]])

section_chip(s2, RX, Inches(3.85), RW, "算例：典型绕组匝间短路")
add_rect(s2, RX, Inches(4.35), RW, Inches(2.55), WHITE, line=BLUE, round_=True)
add_text(s2, RX + Inches(0.2), Inches(4.5), RW - Inches(0.4), Inches(2.3), [
    [R("输入 DGA(μL/L)：", 13, NAVY, True, BODY_FONT), R("H₂=620, CH₄=180,", 12.5, DARK, False, MONO_FONT)],
    [R("                C₂H₂=240, C₂H₄=300, C₂H₆=85", 12.5, DARK, False, MONO_FONT)],
    [R("① 编码 [1,0,2] → 命中“电弧放电”规则(0.93)", 12.5, DARK, False, BODY_FONT)],
    [R("② 多征兆超标，短路 CPT 似然高(0.85~0.90)", 12.5, DARK, False, BODY_FONT)],
    [R("③ 融合：0.7·Pbayes + 0.3·0.93 → 概率抬升", 12.5, DARK, False, BODY_FONT)],
    [R("④ 输出：", 13, NAVY, True, BODY_FONT), R("绕组匝间短路 排名第1【高风险】", 13, ACCENT, True, BODY_FONT)],
], space_after=6, line_spacing=1.05)

prs.save(str(OUT))
print("已生成:", OUT)
