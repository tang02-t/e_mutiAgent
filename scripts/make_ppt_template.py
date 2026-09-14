#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
基于 ts-研究进展汇报-0615.pptx 的模板风格，生成"混合故障归因"内容 PPT。
风格复刻：紫色主题 #7E4381 + 微软雅黑 + 卡片(F9FAFB/F0F7FF) + 标题栏 + 右上页码。
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "混合故障归因_汇报版.pptx"

# ── 模板配色 ──
PURPLE  = RGBColor(0x7E, 0x43, 0x81)   # 主紫色
PURPLE_D= RGBColor(0x5E, 0x30, 0x61)   # 深紫
T1      = RGBColor(0x21, 0x25, 0x29)   # 主文字
T2      = RGBColor(0x49, 0x50, 0x57)   # 次文字
T3      = RGBColor(0x6C, 0x75, 0x7D)   # 弱文字
CARD1   = RGBColor(0xF9, 0xFA, 0xFB)   # 浅灰卡片
CARD2   = RGBColor(0xF0, 0xF7, 0xFF)   # 浅蓝卡片
PURPLE_L= RGBColor(0xF3, 0xEC, 0xF4)   # 浅紫卡片
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
LINEC   = RGBColor(0xE9, 0xEC, 0xEF)

F = "微软雅黑"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def rect(slide, x, y, w, h, fill, line=None, round_=False, lw=1.0):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    return shp


def oval(slide, x, y, d, fill):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.fill.background(); shp.shadow.inherit = False
    return shp


def txt(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        sa=5, ls=1.08):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(2); tf.margin_top = tf.margin_bottom = Pt(2)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(sa); p.line_spacing = ls
        for (t, sz, c, b, it) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.color.rgb = c; r.font.bold = b
            r.font.name = F; r.font.italic = it
    return tb


def R(t, sz=13, c=T1, b=False, it=False):
    return (t, sz, c, b, it)


def header(slide, title, page):
    # 顶部紫色细条 + 标题
    rect(slide, 0, 0, SW, Inches(0.12), PURPLE)
    txt(slide, Inches(0.8), Inches(0.42), Inches(11.0), Inches(0.6),
        [[R(title, 24, PURPLE, True)]], anchor=MSO_ANCHOR.MIDDLE)
    # 标题下分隔线
    rect(slide, Inches(0.8), Inches(1.12), Inches(11.73), Pt(1.5), LINEC)
    # 右上角页码
    txt(slide, SW - Inches(1.6), Inches(0.45), Inches(1.0), Inches(0.5),
        [[R(page, 14, T3, False)]], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def sub_chip(slide, x, y, w, text):
    """小节标题：左侧紫色竖条 + 文字。"""
    rect(slide, x, y + Inches(0.02), Pt(4), Inches(0.32), PURPLE)
    txt(slide, x + Inches(0.16), y, w - Inches(0.2), Inches(0.38),
        [[R(text, 15, T1, True)]], anchor=MSO_ANCHOR.MIDDLE)


# ════════════════════════════════════════════════
# 第 1 页：原理
# ════════════════════════════════════════════════
s1 = prs.slides.add_slide(BLANK)
rect(s1, 0, 0, SW, SH, WHITE)
header(s1, "故障归因方法：三比值法 × 贝叶斯网络混合推理", "/ 原理")

# 顶部一句话定位（浅紫底）
rect(s1, Inches(0.8), Inches(1.35), Inches(11.73), Inches(0.7), PURPLE_L, round_=True)
txt(s1, Inches(1.05), Inches(1.35), Inches(11.3), Inches(0.7),
    [[R("核心思想：", 14, PURPLE, True), R("以贝叶斯概率推理为主、国标三比值规则为辅，融合二者优势——既可量化故障概率，又符合 DL/T 722 行业规范、具备可解释性。", 14, T2)]],
    anchor=MSO_ANCHOR.MIDDLE)

# 左栏：设计动机（对比）
LX = Inches(0.8); LW = Inches(5.6)
sub_chip(s1, LX, Inches(2.3), LW, "为什么要“混合”？")
# 两张对比卡
rect(s1, LX, Inches(2.8), LW, Inches(1.1), CARD1, round_=True)
txt(s1, LX + Inches(0.2), Inches(2.92), LW - Inches(0.4), Inches(0.9), [
    [R("三比值法（DL/T 722 国标）", 14, PURPLE, True)],
    [R("✓ 可解释、合规　　✗ 仅判定类型、无法量化概率", 12.5, T2)],
], sa=3)
rect(s1, LX, Inches(4.0), LW, Inches(1.1), CARD1, round_=True)
txt(s1, LX + Inches(0.2), Inches(4.12), LW - Inches(0.4), Inches(0.9), [
    [R("贝叶斯网络（概率推理）", 14, PURPLE, True)],
    [R("✓ 量化故障概率　　✗ 黑箱、可解释性差", 12.5, T2)],
], sa=3)

sub_chip(s1, LX, Inches(5.35), LW, "贝叶斯网络结构（两层有向图）")
rect(s1, LX, Inches(5.85), Inches(2.5), Inches(1.0), CARD2, line=PURPLE, round_=True, lw=1.2)
txt(s1, LX, Inches(5.85), Inches(2.5), Inches(1.0),
    [[R("故障层 F", 14, PURPLE, True)], [R("8 类故障", 11.5, T3)], [R("短路/局放/过热…", 10.5, T3)]],
    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, sa=1)
ar = s1.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, LX + Inches(2.55), Inches(6.15), Inches(0.5), Inches(0.35))
ar.fill.solid(); ar.fill.fore_color.rgb = PURPLE; ar.line.fill.background(); ar.shadow.inherit = False
rect(s1, LX + Inches(3.1), Inches(5.85), Inches(2.5), Inches(1.0), CARD2, line=PURPLE, round_=True, lw=1.2)
txt(s1, LX + Inches(3.1), Inches(5.85), Inches(2.5), Inches(1.0),
    [[R("征兆层 S", 14, PURPLE, True)], [R("14 类征兆", 11.5, T3)], [R("气体超标/油温升高…", 10.5, T3)]],
    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, sa=1)

# 右栏：四步流程
RX = Inches(6.75); RW = Inches(5.78)
sub_chip(s1, RX, Inches(2.3), RW, "四步推理流程")
steps = [
    ("1", "DGA 三比值法编码", "三个气体比值离散编码为 0/1/2，匹配表7故障判据，得“规则证据”"),
    ("2", "朴素贝叶斯后验推理", "由先验 P(F) 与条件概率表 CPT，计算各故障后验概率并归一化"),
    ("3", "规则—概率融合", "0.7×贝叶斯 + 0.3×规则置信，对命中故障加权修正"),
    ("4", "排序 · 分级 · 报告", "按融合概率排序，映射高/中/低风险，输出可解释报告"),
]
sy = Inches(2.82)
for num, head, desc in steps:
    rect(s1, RX, sy, RW, Inches(0.98), CARD1, round_=True)
    oval(s1, RX + Inches(0.16), sy + Inches(0.23), Inches(0.52), PURPLE)
    txt(s1, RX + Inches(0.16), sy + Inches(0.23), Inches(0.52), Inches(0.52),
        [[R(num, 18, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s1, RX + Inches(0.85), sy + Inches(0.1), RW - Inches(1.0), Inches(0.82), [
        [R(head, 14.5, T1, True)], [R(desc, 11.5, T2)],
    ], sa=2)
    sy += Inches(1.06)

# ════════════════════════════════════════════════
# 第 2 页：公式与算例
# ════════════════════════════════════════════════
s2 = prs.slides.add_slide(BLANK)
rect(s2, 0, 0, SW, SH, WHITE)
header(s2, "故障归因方法：核心公式与诊断算例", "/ 公式")

# 左栏
LX = Inches(0.8); LW = Inches(5.85)
sub_chip(s2, LX, Inches(1.4), LW, "① 三比值法编码规则（DL/T 722 表6）")
tbl = s2.shapes.add_table(4, 2, LX, Inches(1.9), LW, Inches(1.55)).table
tbl.columns[0].width = Inches(2.9); tbl.columns[1].width = Inches(2.95)
data = [("特征比值", "编码规则"),
        ("C₂H₂ / C₂H₄", "<0.1→0 ; [0.1,3)→1 ; ≥3→2"),
        ("CH₄ / H₂", "<0.1→1 ; [0.1,1)→0 ; ≥1→2"),
        ("C₂H₄ / C₂H₆", "<1→0 ; [1,3)→1 ; ≥3→2")]
for i, (a, b) in enumerate(data):
    for j, t in enumerate((a, b)):
        cell = tbl.cell(i, j); cell.fill.solid()
        cell.fill.fore_color.rgb = PURPLE if i == 0 else (CARD1 if i % 2 else WHITE)
        tf = cell.text_frame; tf.word_wrap = True
        tf.margin_top = Pt(3); tf.margin_bottom = Pt(3)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = t
        r.font.size = Pt(12.5 if i == 0 else 12); r.font.bold = (i == 0)
        r.font.color.rgb = WHITE if i == 0 else T1; r.font.name = F

sub_chip(s2, LX, Inches(3.7), LW, "② 朴素贝叶斯后验推理")
rect(s2, LX, Inches(4.2), LW, Inches(0.9), PURPLE_L, round_=True)
txt(s2, LX, Inches(4.2), LW, Inches(0.9),
    [[R("P(Fᵢ | E) ∝ P(Fᵢ) · ∏ P(eⱼ | Fᵢ)", 19, PURPLE_D, True)]],
    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
txt(s2, LX, Inches(5.2), LW, Inches(1.5), [
    [R("• P(Fᵢ) ", 12.5, PURPLE, True), R("故障先验（常见程度，如绝缘老化 0.18）", 12, T2)],
    [R("• P(eⱼ|Fᵢ) ", 12.5, PURPLE, True), R("条件概率表 CPT（故障发生时征兆出现概率）", 12, T2)],
    [R("• 仅对观测到的征兆连乘，最后归一化为概率分布", 12, T2)],
], sa=5, ls=1.12)

# 右栏
RX = Inches(6.95); RW = Inches(5.6)
sub_chip(s2, RX, Inches(1.4), RW, "③ 规则—概率融合")
rect(s2, RX, Inches(1.9), RW, Inches(1.2), PURPLE_L, round_=True)
txt(s2, RX, Inches(1.98), RW, Inches(1.1), [
    [R("P_fused(Fᵢ) =", 15, PURPLE_D, True)],
    [R("0.7·P_bayes + 0.3·P_rule  ", 16, PURPLE_D, True), R("（命中规则）", 11, T3)],
    [R("P_bayes  ", 15, PURPLE_D, True), R("（未命中规则）", 11, T3)],
], align=PP_ALIGN.CENTER, sa=2)
txt(s2, RX, Inches(3.15), RW, Inches(0.4),
    [[R("以数据驱动概率为主(70%)，国标规则经验修正(30%)", 11.5, T3, False, True)]])

sub_chip(s2, RX, Inches(3.7), RW, "诊断算例：典型绕组匝间短路")
rect(s2, RX, Inches(4.2), RW, Inches(2.55), CARD2, line=PURPLE, round_=True, lw=1.2)
txt(s2, RX + Inches(0.22), Inches(4.35), RW - Inches(0.44), Inches(2.3), [
    [R("输入 DGA(μL/L)：", 12.5, PURPLE, True), R("H₂=620, CH₄=180,", 12, T1)],
    [R("                C₂H₂=240, C₂H₄=300, C₂H₆=85", 12, T1)],
    [R("① 编码 [1,0,2] → 命中“电弧放电”规则(0.93)", 12, T2)],
    [R("② 多征兆超标，短路 CPT 似然高(0.85~0.90)", 12, T2)],
    [R("③ 融合：0.7·P_bayes + 0.3·0.93 → 概率抬升", 12, T2)],
    [R("④ 输出：", 12.5, PURPLE, True), R("绕组匝间短路 排名第1 【高风险】", 12.5, PURPLE_D, True)],
], sa=6, ls=1.06)

prs.save(str(OUT))
print("已生成:", OUT)
