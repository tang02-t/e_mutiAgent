#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
将"图3-1 系统总体架构图"生成为 PPT 原生可编辑图形。
所有方框、文字、连线均为 PowerPoint 自带形状，可在 PPT 内自由编辑。
颜色填充与边线虚实严格对应原图：
  - 编排层：浅蓝填充 + 深蓝虚线容器
  - 智能体层：浅绿填充 + 绿色虚线容器
  - 工具层：浅橙填充 + 橙色虚线容器
  - 节点框：圆角矩形 + 实线边框
  - 连线：顺序流(绿实线)/REVISION迭代(绿虚线)/工具调度(橙实线)
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "figures" / "fig3_1_architecture.pptx"
OUT.parent.mkdir(parents=True, exist_ok=True)

# ── 配色 ──
BLUE = RGBColor(0x1F, 0x3A, 0x5F)
LBLUE = RGBColor(0xDC, 0xE6, 0xF1)
BLUE_BORDER = RGBColor(0x9D, 0xB4, 0xCE)
GREEN = RGBColor(0x2E, 0x7D, 0x54)
LGREEN = RGBColor(0xD9, 0xEA, 0xD3)
GREEN_BORDER = RGBColor(0x8F, 0xC0, 0xA2)
ORANGE = RGBColor(0xB5, 0x65, 0x1D)
LORANGE = RGBColor(0xFC, 0xE5, 0xCD)
ORANGE_BORDER = RGBColor(0xD2, 0xA6, 0x79)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x99, 0x99, 0x99)
PANEL_BLUE = RGBColor(0xF2, 0xF6, 0xFB)
PANEL_GREEN = RGBColor(0xF4, 0xFB, 0xF6)
PANEL_ORANGE = RGBColor(0xFD, 0xF7, 0xF0)


def set_dash(line_format, dash="dash"):
    """设置线条虚实：dash / solid。"""
    ln = line_format._get_or_add_ln()
    # 移除已有 prstDash
    for el in ln.findall(qn("a:prstDash")):
        ln.remove(el)
    if dash != "solid":
        d = ln.makeelement(qn("a:prstDash"), {"val": dash})
        ln.append(d)


def add_box(slide, x, y, w, h, text, fill, border, font_color=DARK,
            fontsize=11, bold=True, dash="solid", rounded=True, anchor=MSO_ANCHOR.MIDDLE,
            border_w=1.5):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    sp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    sp.line.color.rgb = border
    sp.line.width = Pt(border_w)
    set_dash(sp.line, dash)
    sp.shadow.inherit = False
    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    for i, line in enumerate(text.split("\n")):
        run = p.add_run() if i == 0 else tf.add_paragraph().add_run()
        if i > 0:
            tf.paragraphs[i].alignment = PP_ALIGN.CENTER
        run.text = line
        run.font.size = Pt(fontsize if i == 0 else fontsize - 1.5)
        run.font.bold = bold
        run.font.color.rgb = font_color
        run.font.name = "微软雅黑"
    return sp


def add_panel(slide, x, y, w, h, title, fill, border, title_color):
    """分层容器：浅色填充 + 虚线边框，标题贴左上。"""
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    sp.line.color.rgb = border
    sp.line.width = Pt(1.25)
    set_dash(sp.line, "dash")
    sp.shadow.inherit = False
    sp.text_frame.text = ""
    # 标题文字框（左上角）
    tb = slide.shapes.add_textbox(Inches(x + 0.1), Inches(y + 0.05), Inches(1.5), Inches(0.3))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = title
    r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = title_color
    r.font.name = "微软雅黑"
    return sp


def add_connector(slide, x1, y1, x2, y2, color, dash="solid", width=1.5):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                    Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    cn.line.color.rgb = color
    cn.line.width = Pt(width)
    set_dash(cn.line, dash)
    # 末端箭头
    ln = cn.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return cn


def _remove_arrow(connector):
    """移除连接器末端箭头（用于多段折线的中间段）。"""
    ln = connector.line._get_or_add_ln()
    for tail in ln.findall(qn("a:tailEnd")):
        ln.remove(tail)


def add_label(slide, x, y, text, color, size=9):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(2.0), Inches(0.3))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = text
    r.font.size = Pt(size); r.font.italic = True; r.font.color.rgb = color
    r.font.name = "微软雅黑"


def main():
    prs = Presentation()
    prs.slide_width = Inches(12)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白版式

    # 标题
    tb = slide.shapes.add_textbox(Inches(0), Inches(0.15), Inches(12), Inches(0.5))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = "图 3-1  系统总体架构图"
    r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = DARK
    r.font.name = "微软雅黑"
    tb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    # ── 三个分层容器 ──
    add_panel(slide, 0.4, 0.9, 11.2, 1.3, "编排层", PANEL_BLUE, BLUE_BORDER, BLUE)
    add_panel(slide, 0.4, 2.5, 11.2, 2.0, "智能体层", PANEL_GREEN, GREEN_BORDER, GREEN)
    add_panel(slide, 0.4, 5.1, 11.2, 1.9, "工具层", PANEL_ORANGE, ORANGE_BORDER, ORANGE)

    # ── 编排层节点 ──
    add_box(slide, 3.0, 1.3, 6.0, 0.7,
            "LangGraph 图式编排\n有向状态图 / 条件路由 / 共享状态 AgentState",
            LBLUE, BLUE, fontsize=11)

    # ── 智能体层 四节点 ──
    ag_x = [0.8, 3.6, 6.4, 9.2]
    ag_txt = ["Planner\n意图分析+任务分解", "Retriever\n工具调度(并行/串行)",
              "Generator\n诊断结论生成", "Validator\n质量验证+路由"]
    ag_w, ag_h, ag_y = 2.2, 1.0, 3.1
    for x, t in zip(ag_x, ag_txt):
        add_box(slide, x, ag_y, ag_w, ag_h, t, LGREEN, GREEN, fontsize=11)

    # ── 工具层 四节点 ──
    tl_x = [0.8, 3.6, 6.4, 9.2]
    tl_txt = ["RAG\n知识检索", "贝叶斯\n故障归因", "ETT\n油温预测", "时序\n异常检测"]
    tl_w, tl_h, tl_y = 2.2, 0.9, 5.6
    for x, t in zip(tl_x, tl_txt):
        add_box(slide, x, tl_y, tl_w, tl_h, t, LORANGE, ORANGE, fontsize=11)

    # ── 连线 ──
    cy = ag_y + ag_h / 2  # 智能体中线
    # 编排层 -> Planner（蓝实线）
    add_connector(slide, 6.0, 2.0, ag_x[0] + ag_w / 2, ag_y, BLUE, "solid")
    # 智能体顺序流（绿实线）
    for i in range(3):
        add_connector(slide, ag_x[i] + ag_w, cy, ag_x[i + 1], cy, GREEN, "solid")
    # Validator -> Generator 迭代回退（绿虚线，走节点下方更明显）
    back_y = ag_y + ag_h + 0.22
    vx = ag_x[3] + ag_w / 2   # Validator 中心x
    gx = ag_x[2] + ag_w / 2   # Generator 中心x
    # 下行段（无箭头）
    cn1 = add_connector(slide, vx, ag_y + ag_h, vx, back_y, GREEN, "dash")
    _remove_arrow(cn1)
    # 水平回退段（无箭头）
    cn2 = add_connector(slide, vx, back_y, gx, back_y, GREEN, "dash")
    _remove_arrow(cn2)
    # 上行段（带箭头指回 Generator）
    add_connector(slide, gx, back_y, gx, ag_y + ag_h, GREEN, "dash")
    add_label(slide, (gx + vx) / 2 - 0.55, back_y - 0.28, "REVISION 迭代", GREEN)
    # Retriever -> 工具层四节点（橙实线）
    for x in tl_x:
        add_connector(slide, ag_x[1] + ag_w / 2, ag_y + ag_h,
                      x + tl_w / 2, tl_y, ORANGE, "solid", width=1.2)

    prs.save(OUT)
    print("已生成 PPT:", OUT)


if __name__ == "__main__":
    main()
